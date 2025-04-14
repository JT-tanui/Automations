import os
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
from PIL import Image
import pdf2image
import tempfile
from pptx import Presentation
from pptx.util import Inches
import img2pdf  # Add this import for image to PDF conversion

class PDFToPPTConverter:
    def __init__(self, root=None):
        self.pdf_files = []
        self.image_files = []
        self.temp_dir = tempfile.mkdtemp()
        self.poppler_path = r"C:\tools\poppler-24.08.0\Library\bin"
        
        if root:
            self.root = root
        else:
            self.root = tk.Tk()
            self.root.title("Document Converter")
            self.root.geometry("600x500")
            self.setup_ui()
    
    def setup_ui(self):
        # Create a notebook with tabs
        self.notebook = ttk.Notebook(self.root)
        self.notebook.pack(fill=tk.BOTH, expand=True)
        
        # Create PPT tab
        self.ppt_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.ppt_tab, text="PDF to PowerPoint")
        
        # Create PDF tab
        self.pdf_tab = ttk.Frame(self.notebook)
        self.notebook.add(self.pdf_tab, text="Images to PDF")
        
        # Set up PPT tab (existing functionality)
        self.setup_ppt_tab()
        
        # Set up PDF tab (new functionality)
        self.setup_pdf_tab()
    
    def setup_ppt_tab(self):
        # Create frames
        self.top_frame = tk.Frame(self.ppt_tab)
        self.top_frame.pack(fill=tk.X, padx=10, pady=10)
        
        self.list_frame = tk.Frame(self.ppt_tab)
        self.list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        self.button_frame = tk.Frame(self.ppt_tab)
        self.button_frame.pack(fill=tk.X, padx=10, pady=10)
        
        # Top frame buttons
        self.select_folder_btn = tk.Button(self.top_frame, text="Select Default Folder", command=self.select_default_folder)
        self.select_folder_btn.pack(side=tk.LEFT, padx=5)
        
        self.select_files_btn = tk.Button(self.top_frame, text="Select Files", command=self.select_pdf_files)
        self.select_files_btn.pack(side=tk.LEFT, padx=5)
        
        # File List
        self.list_label = tk.Label(self.list_frame, text="Selected Files (PDFs and Images):")
        self.list_label.pack(anchor=tk.W)
        
        self.listbox_frame = tk.Frame(self.list_frame)
        self.listbox_frame.pack(fill=tk.BOTH, expand=True)
        
        self.scrollbar = tk.Scrollbar(self.listbox_frame)
        self.scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.pdf_listbox = tk.Listbox(self.listbox_frame)
        self.pdf_listbox.pack(fill=tk.BOTH, expand=True)
        
        self.pdf_listbox.config(yscrollcommand=self.scrollbar.set)
        self.scrollbar.config(command=self.pdf_listbox.yview)
        
        # Button frame
        self.remove_btn = tk.Button(self.button_frame, text="Remove Selected", command=self.remove_selected)
        self.remove_btn.pack(side=tk.LEFT, padx=5)
        
        self.clear_btn = tk.Button(self.button_frame, text="Clear All", command=self.clear_all)
        self.clear_btn.pack(side=tk.LEFT, padx=5)
        
        self.generate_btn = tk.Button(self.button_frame, text="Generate PowerPoint", command=self.generate_ppt)
        self.generate_btn.pack(side=tk.RIGHT, padx=5)
    
    def setup_pdf_tab(self):
        # Create frames
        self.pdf_top_frame = tk.Frame(self.pdf_tab)
        self.pdf_top_frame.pack(fill=tk.X, padx=10, pady=10)
        
        self.pdf_list_frame = tk.Frame(self.pdf_tab)
        self.pdf_list_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
        
        self.pdf_button_frame = tk.Frame(self.pdf_tab)
        self.pdf_button_frame.pack(fill=tk.X, padx=10, pady=10)
        
        # Top frame buttons
        self.select_images_btn = tk.Button(self.pdf_top_frame, text="Select Images", command=self.select_images)
        self.select_images_btn.pack(side=tk.LEFT, padx=5)
        
        self.select_images_folder_btn = tk.Button(self.pdf_top_frame, text="Select Folder", command=self.select_images_folder)
        self.select_images_folder_btn.pack(side=tk.LEFT, padx=5)
        
        # File List
        self.image_list_label = tk.Label(self.pdf_list_frame, text="Selected Images:")
        self.image_list_label.pack(anchor=tk.W)
        
        self.image_listbox_frame = tk.Frame(self.pdf_list_frame)
        self.image_listbox_frame.pack(fill=tk.BOTH, expand=True)
        
        self.image_scrollbar = tk.Scrollbar(self.image_listbox_frame)
        self.image_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.image_listbox = tk.Listbox(self.image_listbox_frame)
        self.image_listbox.pack(fill=tk.BOTH, expand=True)
        
        self.image_listbox.config(yscrollcommand=self.image_scrollbar.set)
        self.image_scrollbar.config(command=self.image_listbox.yview)
        
        # Button frame
        self.remove_image_btn = tk.Button(self.pdf_button_frame, text="Remove Selected", command=self.remove_selected_image)
        self.remove_image_btn.pack(side=tk.LEFT, padx=5)
        
        self.clear_images_btn = tk.Button(self.pdf_button_frame, text="Clear All", command=self.clear_all_images)
        self.clear_images_btn.pack(side=tk.LEFT, padx=5)
        
        self.generate_pdf_btn = tk.Button(self.pdf_button_frame, text="Generate PDF", command=self.generate_pdf)
        self.generate_pdf_btn.pack(side=tk.RIGHT, padx=5)
    
    def select_default_folder(self):
        """Look for PDFs in the 'cert' folder in the root project directory."""
        cert_folder = os.path.join(os.getcwd(), "cert")
        
        if os.path.exists(cert_folder) and os.path.isdir(cert_folder):
            self.add_pdfs_from_folder(cert_folder)
        else:
            messagebox.showinfo("Info", "The 'cert' folder does not exist. Please create it or select files manually.")
            folder_path = filedialog.askdirectory(title="Select Folder Containing PDFs and Images")
            if folder_path:
                self.add_pdfs_from_folder(folder_path)
    
    def add_pdfs_from_folder(self, folder_path):
        """Add all supported files from the specified folder to the list."""
        # Define allowed extensions
        allowed_extensions = ('.pdf', '.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff')
        
        # Find all files with allowed extensions
        new_files = [os.path.join(folder_path, f) for f in os.listdir(folder_path) 
                    if os.path.splitext(f.lower())[1] in allowed_extensions]
        
        if not new_files:
            messagebox.showinfo("Info", "No supported files found in the selected folder.")
            return
        
        # Add unique files
        for file_path in new_files:
            if file_path not in self.pdf_files:
                self.pdf_files.append(file_path)
                self.pdf_listbox.insert(tk.END, os.path.basename(file_path))
    
    def select_pdf_files(self):
        """Open file dialog to select PDF files and images."""
        file_paths = filedialog.askopenfilenames(
            title="Select Files",
            filetypes=[
                ("All Supported Files", "*.pdf *.jpg *.jpeg *.png *.bmp *.gif *.tiff"),
                ("PDF Files", "*.pdf"),
                ("Image Files", "*.jpg *.jpeg *.png *.bmp *.gif *.tiff"),
            ]
        )
        
        if not file_paths:
            return
        
        # Add unique files
        for file_path in file_paths:
            if file_path not in self.pdf_files:
                self.pdf_files.append(file_path)
                self.pdf_listbox.insert(tk.END, os.path.basename(file_path))
    
    def remove_selected(self):
        """Remove selected PDFs from the list."""
        selection = self.pdf_listbox.curselection()
        if not selection:
            messagebox.showinfo("Info", "Please select a PDF to remove.")
            return
        
        # Remove from bottom to top to maintain indices
        for index in sorted(selection, reverse=True):
            del self.pdf_files[index]
            self.pdf_listbox.delete(index)
    
    def clear_all(self):
        """Clear all PDFs from the list."""
        self.pdf_files = []
        self.pdf_listbox.delete(0, tk.END)
    
    def generate_ppt(self):
        """Generate PowerPoint from selected PDFs and images."""
        if not self.pdf_files:
            messagebox.showinfo("Info", "Please add at least one file.")
            return
        
        # Ask for output filename
        output_file = filedialog.asksaveasfilename(
            title="Save PowerPoint As",
            defaultextension=".pptx",
            filetypes=[("PowerPoint Presentation", "*.pptx")]
        )
        
        if not output_file:
            return
        
        try:
            # Create presentation
            prs = Presentation()
            
            # Process each file
            for file_path in self.pdf_files:
                file_extension = os.path.splitext(file_path)[1].lower()
                
                # Handle based on file type
                if file_extension == '.pdf':
                    self.add_pdf_to_ppt(prs, file_path)
                elif file_extension in ['.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff']:
                    self.add_image_to_ppt(prs, file_path)
                else:
                    print(f"Unsupported file type: {file_extension}")
            
            # Save presentation
            prs.save(output_file)
            messagebox.showinfo("Success", f"PowerPoint created successfully!\nSaved to: {output_file}")
        
        except Exception as e:
            messagebox.showerror("Error", f"An error occurred: {str(e)}")
    
    def add_pdf_to_ppt(self, prs, pdf_path):
        """Convert a PDF page to an image and add it to the presentation."""
        try:
            # Convert PDF to images with higher quality
            images = pdf2image.convert_from_path(
                pdf_path,
                poppler_path=self.poppler_path,
                dpi=300  # Increased DPI for better quality
            )
            
            for img in images:
                # Save image temporarily with better quality
                img_path = os.path.join(self.temp_dir, f"temp_img_{os.path.basename(pdf_path)}_{id(img)}.png")
                img.save(img_path, "PNG", optimize=True)
                
                # Add slide
                slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
                
                # Get slide dimensions
                slide_width = prs.slide_width
                slide_height = prs.slide_height
                
                # Add image to slide - starting position doesn't matter, will be adjusted
                pic = slide.shapes.add_picture(img_path, 0, 0)
                
                # Apply small margin (2.5%) for better visual appearance
                margin_factor = 0.95  # 95% of available space
                max_width = slide_width * margin_factor
                max_height = slide_height * margin_factor
                
                # Calculate scaling to preserve aspect ratio
                width_scale = max_width / pic.width
                height_scale = max_height / pic.height
                scale = min(width_scale, height_scale)
                
                # Apply scaling
                new_width = int(pic.width * scale)
                new_height = int(pic.height * scale)
                pic.width = new_width
                pic.height = new_height
                
                # Center precisely on slide
                pic.left = int((slide_width - new_width) / 2)
                pic.top = int((slide_height - new_height) / 2)
                
                # Clean up temp file
                os.remove(img_path)
                
                # Report progress
                print(f"Added slide from {os.path.basename(pdf_path)}")
                
        except Exception as e:
            error_msg = f"Error processing {os.path.basename(pdf_path)}: {str(e)}"
            print(error_msg)
            messagebox.showerror("Processing Error", error_msg)
            raise
    
    def add_image_to_ppt(self, prs, image_path):
        """Add an image directly to the presentation."""
        try:
            # Add slide
            slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
            
            # Get slide dimensions
            slide_width = prs.slide_width
            slide_height = prs.slide_height
            
            # Add image to slide
            pic = slide.shapes.add_picture(image_path, 0, 0)
            
            # Apply small margin for better visual appearance
            margin_factor = 0.95  # 95% of available space
            max_width = slide_width * margin_factor
            max_height = slide_height * margin_factor
            
            # Calculate scaling to preserve aspect ratio
            width_scale = max_width / pic.width
            height_scale = max_height / pic.height
            scale = min(width_scale, height_scale)
            
            # Apply scaling
            new_width = int(pic.width * scale)
            new_height = int(pic.height * scale)
            pic.width = new_width
            pic.height = new_height
            
            # Center precisely on slide
            pic.left = int((slide_width - new_width) / 2)
            pic.top = int((slide_height - new_height) / 2)
            
            # Report progress
            print(f"Added slide from {os.path.basename(image_path)}")
                
        except Exception as e:
            error_msg = f"Error processing {os.path.basename(image_path)}: {str(e)}"
            print(error_msg)
            messagebox.showerror("Processing Error", error_msg)
            raise
    
    def select_images(self):
        """Open file dialog to select image files."""
        file_paths = filedialog.askopenfilenames(
            title="Select Images",
            filetypes=[
                ("Image Files", "*.jpg *.jpeg *.png *.bmp *.gif *.tiff"),
            ]
        )
        
        if not file_paths:
            return
        
        # Add unique files
        for file_path in file_paths:
            if file_path not in self.image_files:
                self.image_files.append(file_path)
                self.image_listbox.insert(tk.END, os.path.basename(file_path))
    
    def select_images_folder(self):
        """Select folder containing images."""
        folder_path = filedialog.askdirectory(title="Select Folder Containing Images")
        if not folder_path:
            return
        
        # Define allowed image extensions
        allowed_extensions = ('.jpg', '.jpeg', '.png', '.bmp', '.gif', '.tiff')
        
        # Find all image files
        new_files = [os.path.join(folder_path, f) for f in os.listdir(folder_path) 
                    if os.path.splitext(f.lower())[1] in allowed_extensions]
        
        if not new_files:
            messagebox.showinfo("Info", "No image files found in the selected folder.")
            return
        
        # Add unique files
        for file_path in new_files:
            if file_path not in self.image_files:
                self.image_files.append(file_path)
                self.image_listbox.insert(tk.END, os.path.basename(file_path))
    
    def remove_selected_image(self):
        """Remove selected images from the list."""
        selection = self.image_listbox.curselection()
        if not selection:
            messagebox.showinfo("Info", "Please select an image to remove.")
            return
        
        # Remove from bottom to top to maintain indices
        for index in sorted(selection, reverse=True):
            del self.image_files[index]
            self.image_listbox.delete(index)
    
    def clear_all_images(self):
        """Clear all images from the list."""
        self.image_files = []
        self.image_listbox.delete(0, tk.END)
    
    def generate_pdf(self):
        """Generate PDF from selected images."""
        if not self.image_files:
            messagebox.showinfo("Info", "Please add at least one image.")
            return
        
        # Ask for output filename
        output_file = filedialog.asksaveasfilename(
            title="Save PDF As",
            defaultextension=".pdf",
            filetypes=[("PDF Document", "*.pdf")]
        )
        
        if not output_file:
            return
        
        try:
            # Convert images to PDF
            with open(output_file, "wb") as f:
                f.write(img2pdf.convert([file for file in self.image_files]))
            
            messagebox.showinfo("Success", f"PDF created successfully!\nSaved to: {output_file}")
        
        except Exception as e:
            messagebox.showerror("Error", f"An error occurred: {str(e)}")

    def run(self):
        """Start the application."""
        self.root.mainloop()

if __name__ == "__main__":
    converter = PDFToPPTConverter()
    converter.run()